from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


REQUIRED_HEADERS = (
    "DATA", "NOME", "ENTRADA 1", "SAÍDA 1", "ENTRADA 2", "SAÍDA 2",
    "ENTRADA 3", "SAÍDA 3", "ENTRADA 4", "SAÍDA 4", "ENTRADA 5",
    "SAÍDA 5", "departamento", "re",
)


class InvalidWorkbook(ValueError):
    pass


@dataclass(frozen=True)
class AbsRow:
    work_date: date
    matricula: str
    source_name: str
    source_department: str
    punches: tuple[str, ...]
    raw_status: str | None
    validated_status: str


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _marker(value: str) -> str:
    return "".join(char for char in unicodedata.normalize("NFD", value.upper())
                   if unicodedata.category(char) != "Mn").strip()


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(value))
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(_text(value), pattern).date()
        except ValueError:
            pass
    raise InvalidWorkbook(f"DATA inválida: {value!r}")


def _classify(punches: tuple[str, ...]) -> tuple[str | None, str]:
    if any(re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2})?", value) for value in punches):
        return None, "P"
    markers = {
        "P": "P", "ATES": "FJ", "DECLARA": "FJ", "FALTA": "FI",
        "FERIAS": "FE", "INSS": "INSS", "L. MATE": "LM", "FOLGA": "FG",
        "AMAMENT": "LAM", "LICENCA PA": "LP", "MT": "MATRIMONIO",
    }
    for value in punches:
        if _marker(value) in markers:
            return value, markers[_marker(value)]
    return (punches[0], "-") if punches else (None, "-")


def read_abs_workbook(path: Path) -> list[AbsRow]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as error:
        raise InvalidWorkbook(f"Não foi possível abrir a base: {error}") from error
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise InvalidWorkbook("Aba obrigatória Sheet1 não encontrada.")
        sheet = workbook["Sheet1"]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = tuple(_text(value) for value in header_values)
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            raise InvalidWorkbook(f"Cabeçalhos obrigatórios ausentes: {', '.join(missing)}")
        index = {header: position for position, header in enumerate(headers)}
        rows: list[AbsRow] = []
        punch_headers = [f"{kind} {number}" for number in range(1, 6) for kind in ("ENTRADA", "SAÍDA")]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in values):
                continue
            punches = tuple(filter(None, (_text(values[index[name]]) for name in punch_headers)))
            raw_status, validated_status = _classify(punches)
            rows.append(AbsRow(
                work_date=_date(values[index["DATA"]]),
                matricula=_text(values[index["re"]]),
                source_name=_text(values[index["NOME"]]),
                source_department=_text(values[index["departamento"]]),
                punches=punches,
                raw_status=raw_status,
                validated_status=validated_status,
            ))
        return rows
    finally:
        workbook.close()
