import argparse
import json
import os
from pathlib import Path

from agent_abs.config import load_config
from agent_abs.api import SupabaseAbsApi
from openpyxl import load_workbook

def get_actual_data(api: SupabaseAbsApi) -> dict:
    import json
    from urllib.request import Request, urlopen
    
    url = f"{api.base}/mop_abs_records?work_date=gte.2026-09-01&work_date=lt.2026-10-01&select=matricula,validated_status"
    req = Request(url, headers=api.headers)
    with urlopen(req, timeout=60) as resp:
        rows = json.loads(resp.read())
        
    result = {}
    for row in rows:
        matricula = str(row["matricula"]).zfill(4)
        if matricula not in result:
            result[matricula] = {"fj": 0, "fi": 0, "faltas": 0, "p": 0, "abs": 0.0}
            
        status = row["validated_status"]
        if status in ("FJ", "LM", "LP", "FE", "INSS", "LAM", "MATRIMONIO"):
            result[matricula]["fj"] += 1
        elif status == "FI":
            result[matricula]["fi"] += 1
        elif status == "P":
            result[matricula]["p"] += 1
            
    for matricula, data in result.items():
        faltas = data["fj"] + data["fi"]
        data["faltas"] = faltas
        p = data["p"]
        data["abs"] = faltas / (p + faltas) if (p + faltas) > 0 else 0.0
        
    return result

def get_expected_data(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb["ABS_OPERAÇÃO"]
        result = {}
        for row in sheet.iter_rows(min_row=18, values_only=True):
            if not row[7]:
                continue
            matricula = str(row[7]).zfill(4)
            fj = int(row[12] or 0)
            fi = int(row[13] or 0)
            faltas = int(row[14] or 0)
            p = int(row[15] or 0)
            abs_perc = float(row[16] or 0.0)
            result[matricula] = {"fj": fj, "fi": fi, "faltas": faltas, "p": p, "abs": abs_perc}
        return result
    finally:
        wb.close()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--vr-file", required=True)
    args = parser.parse_args()
    
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).parent / ".env")
    
    environment = dict(os.environ)
    if "ABS_FILE_PATH" not in environment:
        environment["ABS_FILE_PATH"] = "dummy.xlsx"
        
    config = load_config(environment, require_api=True)
    api = SupabaseAbsApi(config.supabase_url, config.service_key)
    
    actual = get_actual_data(api)
    expected = get_expected_data(Path(args.vr_file))
    
    matched = 0
    differences = []
    
    all_keys = set(actual.keys()) | set(expected.keys())
    for k in all_keys:
        if k not in expected:
            differences.append({"matricula": k, "reason": "No MOP mas ausente no VR"})
            continue
        if k not in actual:
            differences.append({"matricula": k, "reason": "No VR mas ausente no MOP (Cadastro divergente)"})
            continue
            
        a = actual[k]
        e = expected[k]
        
        diffs = {}
        for metric in ["fj", "fi", "faltas", "p"]:
            if a[metric] != e[metric]:
                diffs[metric] = {"expected": e[metric], "actual": a[metric]}
                
        if abs(a["abs"] - e["abs"]) > 0.0001:
            diffs["abs"] = {"expected": e["abs"], "actual": a["abs"]}
            
        if diffs:
            # Check if reason is mostly "Declara" -> FJ differences
            if "fj" in diffs and a["fj"] > e["fj"]:
                reason = "Divergência deliberada (Declara -> FJ) ou outras"
            else:
                reason = "Divergência métrica"
            differences.append({"matricula": k, "diffs": diffs, "reason": reason})
        else:
            matched += 1
            
    summary = {
        "matched": matched,
        "expected_total": len(expected),
        "actual_total": len(actual),
        "differences": differences
    }
    print(json.dumps(summary, indent=2, ensure_ascii=False))

if __name__ == "__main__":
    main()
